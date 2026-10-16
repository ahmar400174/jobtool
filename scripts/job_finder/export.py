"""Export job listings to Excel."""

from __future__ import annotations

import re
from datetime import datetime
from html import unescape
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from job_finder.config import project_root
from job_finder.models import JobListing

# Internal field → friendly Excel header
COLUMN_MAP = [
    ("title", "Job Title"),
    ("company", "Company"),
    ("location", "Company Location"),
    ("remote", "Remote"),
    ("posted", "Posted Date"),
    ("source", "Source"),
    ("url", "Apply Link"),
    ("salary", "Salary"),
    ("tags", "Tags"),
    ("description", "Description"),
    ("fetched_at", "Fetched At"),
]

COLUMN_WIDTHS = {
    "Job Title": 38,
    "Company": 26,
    "Company Location": 22,
    "Remote": 10,
    "Posted Date": 14,
    "Source": 12,
    "Apply Link": 55,
    "Salary": 16,
    "Tags": 30,
    "Description": 70,
    "Fetched At": 18,
}


def _strip_html(text: str) -> str:
    if not text:
        return ""
    text = unescape(text)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</p>\s*", "\n", text, flags=re.I)
    text = re.sub(r"<li[^>]*>", "• ", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _excel_text(value: Any) -> str:
    """Public feeds may contain control characters that abort openpyxl writes."""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    # Surrogates and the two XML noncharacters cannot be serialized either.
    return re.sub(r"[\ud800-\udfff\ufffe\uffff]", "", text)[:32767]


def export_excel(jobs: list[JobListing], settings: dict[str, Any]) -> Path:
    out_cfg = settings.get("output", {})
    out_dir = project_root() / out_cfg.get("directory", "output")
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = out_cfg.get("filename", "qa_jobs_usa_remote.xlsx")
    sheet_name = out_cfg.get("sheet_name", "Remote QA Jobs")
    path = out_dir / filename

    rows = []
    for job in jobs:
        row = job.to_row()
        row["description"] = _strip_html(str(row.get("description") or ""))[:2000]
        row["salary"] = str(row.get("salary") or "").strip()
        row["tags"] = str(row.get("tags") or "").strip()
        row = {key: _excel_text(value) for key, value in row.items()}
        rows.append(row)

    internal_cols = [c[0] for c in COLUMN_MAP]
    headers = [c[1] for c in COLUMN_MAP]
    df = pd.DataFrame(rows, columns=internal_cols)
    df.columns = headers
    df = df.fillna("")

    def _write_excel(target: Path) -> None:
        # Own the file handle so it closes even when the writer's save fails.
        with target.open("wb") as stream, pd.ExcelWriter(stream, engine="openpyxl") as writer:
            writer.book.create_sheet(sheet_name)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]

            header_fill = PatternFill(start_color="0A66C2", end_color="0A66C2", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                letter = get_column_letter(col_idx)
                ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(header, 18)

            for row_idx in range(2, len(df) + 2):
                # Prevent Excel from auto-expanding long descriptions into rows
                # that fill most of the screen. Full text remains in the cell.
                ws.row_dimensions[row_idx].height = 30
                for cell in ws[row_idx]:
                    cell.alignment = Alignment(vertical="top")
                    # Listings are external text, including strings starting '='.
                    cell.data_type = "s"
                ws.cell(row=row_idx, column=headers.index("Description") + 1).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
                ws.cell(row=row_idx, column=headers.index("Apply Link") + 1).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

    # Write and verify privately first: failures must never leave a truncated
    # workbook at the public output path or destroy the previous good export.
    with NamedTemporaryFile(dir=out_dir, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        _write_excel(temporary_path)
        workbook = load_workbook(temporary_path, read_only=True)
        try:
            worksheet = workbook[sheet_name]
            saved_rows = sum(1 for _ in worksheet.iter_rows(min_row=2, values_only=True))
            if saved_rows != len(jobs):
                raise RuntimeError(f"Excel export incomplete: expected {len(jobs)} jobs, saved {saved_rows}")
        finally:
            workbook.close()
        try:
            temporary_path.replace(path)
        except PermissionError:
            # An open Excel workbook locks the original. Each fallback gets a
            # fresh name so an earlier '_new' workbook cannot be reopened.
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            path = path.with_name(f"{path.stem}_{stamp}{path.suffix}")
            temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)

    return path
