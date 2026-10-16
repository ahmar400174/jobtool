"""Regression checks for complete exports and reporting the actual saved file."""
import json
import sys
from pathlib import Path
from unittest.mock import patch

import pytest
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from job_finder.aggregator import SourceCounts
from job_finder.export import export_excel
from job_finder.models import JobListing
from job_finder.report import save_search_report


def jobs(count=54):
    return [JobListing(title=f"Software Engineer {i}", company=f"Acme {i}",
                       location="Worldwide", remote="Remote", posted="2026-10-14",
                       url=f"https://example.com/{i}", source="Test",
                       description="<p>Remote engineering role.</p>" * 100) for i in range(count)]


def settings(directory):
    return {"output": {"directory": str(directory), "filename": "jobs.xlsx", "sheet_name": "Jobs"}}


def test_exports_all_54_jobs_even_with_invalid_text_at_row_nine(tmp_path):
    records = jobs()
    records[8].title += "\x00\x0b\ud800\ufffe"
    records[8].description += "&#11;"
    records[8].company = "=External text"
    path = export_excel(records, settings(tmp_path))
    workbook = load_workbook(path)
    try:
        sheet = workbook["Jobs"]
        assert sheet.max_row == 55
        assert sheet["A55"].value == "Software Engineer 53"
        assert sheet["A10"].value == "Software Engineer 8"
        assert sheet["B10"].value == "=External text"
        assert sheet["B10"].data_type == "s"
        assert all(sheet.row_dimensions[row].height == 30 for row in range(2, 56))
        assert not any(dim.hidden for dim in sheet.row_dimensions.values())
        assert sheet.auto_filter.ref == "A1:K55"
    finally:
        workbook.close()


def test_failed_write_preserves_previous_workbook_and_removes_temp_file(tmp_path):
    cfg = settings(tmp_path)
    path = export_excel(jobs(8), cfg)
    original = path.read_bytes()
    write = pd.DataFrame.to_excel
    def fail_after_eight_rows(frame, *args, **kwargs):
        write(frame.iloc[:8], *args, **kwargs)
        raise RuntimeError("write failed after eight rows")
    with patch.object(pd.DataFrame, "to_excel", fail_after_eight_rows):
        with pytest.raises(RuntimeError, match="write failed after eight rows"):
            export_excel(jobs(), cfg)
    assert path.read_bytes() == original
    assert list(tmp_path.iterdir()) == [path]


def test_locked_output_uses_fresh_fallback_each_time(tmp_path):
    cfg = settings(tmp_path)
    original = export_excel(jobs(8), cfg)
    original_bytes = original.read_bytes()
    replace = Path.replace

    def locked(source, target):
        if target == original:
            raise PermissionError("Excel has the file open")
        return replace(source, target)

    with patch.object(Path, "replace", locked):
        first = export_excel(jobs(), cfg)
        second = export_excel(jobs(), cfg)
    assert first != second != original
    assert first.exists() and second.exists()
    assert original.read_bytes() == original_bytes
    workbook = load_workbook(second, read_only=True)
    try:
        assert workbook["Jobs"].max_row == 55
    finally:
        workbook.close()


def test_incomplete_verification_does_not_publish(tmp_path):
    cfg = settings(tmp_path)
    original = export_excel(jobs(8), cfg)
    original_bytes = original.read_bytes()
    actual_load = load_workbook
    def incomplete(path, **kwargs):
        return actual_load(original, **kwargs)
    with patch("job_finder.export.load_workbook", side_effect=incomplete):
        with pytest.raises(RuntimeError, match="expected 54 jobs, saved 8"):
            export_excel(jobs(), cfg)
    assert original.read_bytes() == original_bytes
    assert list(tmp_path.iterdir()) == [original]


def test_gui_exports_and_reports_the_same_54_matches(tmp_path):
    from gui import JobFinderApp
    app = JobFinderApp.__new__(JobFinderApp)
    callbacks, successes = [], []
    app._ui = callbacks.append
    app._log = lambda *args: None
    app._on_success = lambda *args: successes.append(args)
    cfg = settings(tmp_path)
    with patch("job_finder.config.load_settings", return_value=cfg), \
            patch("job_finder.aggregator.fetch_all", return_value=(jobs(), SourceCounts())):
        app._run_fetch()
    for callback in callbacks:
        callback()
    path, count, _ = successes[0]
    assert count == 54
    report = json.loads((tmp_path / "latest_search_report.json").read_text())
    assert report["unique_matches"] == 54
    assert report["excel_export"]["status"] == "saved"
    assert report["excel_export"]["verified_rows"] == 54
    assert report["excel_export"]["path"] == str(path.resolve())


def test_gui_reports_export_failure_separately_from_search(tmp_path):
    from gui import JobFinderApp
    app = JobFinderApp.__new__(JobFinderApp)
    callbacks, errors = [], []
    app._ui = callbacks.append
    app._log = lambda *args: None
    app._on_error = errors.append
    cfg = settings(tmp_path)
    with patch("job_finder.config.load_settings", return_value=cfg), \
            patch("job_finder.aggregator.fetch_all", return_value=(jobs(), SourceCounts())), \
            patch("job_finder.export.export_excel", side_effect=PermissionError("Output folder is read-only")):
        app._run_fetch()
    for callback in callbacks:
        callback()
    assert "Found 54 jobs, but Excel export failed" in errors[0]
    report = json.loads((tmp_path / "latest_search_report.json").read_text())
    assert report["excel_export"]["status"] == "failed"
    assert report["excel_export"]["verified_rows"] is None


def test_empty_search_report_does_not_claim_an_export(tmp_path):
    path = save_search_report(settings(tmp_path), [], SourceCounts())
    report = json.loads(path.read_text())
    assert report["excel_export"]["status"] == "no_matches"


def test_cli_exports_all_matches_and_reports_verified_count(tmp_path, capsys):
    import fetch_jobs
    cfg = settings(tmp_path)
    with patch.object(sys, "argv", ["fetch_jobs.py"]), \
            patch.object(fetch_jobs, "load_settings", return_value=cfg), \
            patch.object(fetch_jobs, "fetch_all", return_value=(jobs(), SourceCounts())):
        assert fetch_jobs.main() == 0
    assert "Excel saved and verified: 54 job rows" in capsys.readouterr().out
    report = json.loads((tmp_path / "latest_search_report.json").read_text())
    assert report["excel_export"]["verified_rows"] == 54


def test_cli_failed_export_returns_failure_and_records_error(tmp_path, capsys):
    import fetch_jobs
    cfg = settings(tmp_path)
    with patch.object(sys, "argv", ["fetch_jobs.py"]), \
            patch.object(fetch_jobs, "load_settings", return_value=cfg), \
            patch.object(fetch_jobs, "fetch_all", return_value=(jobs(), SourceCounts())), \
            patch.object(fetch_jobs, "export_excel", side_effect=RuntimeError("write failed")):
        assert fetch_jobs.main() == 1
    assert "Found 54 jobs, but Excel export failed" in capsys.readouterr().out
    report = json.loads((tmp_path / "latest_search_report.json").read_text())
    assert report["excel_export"]["status"] == "failed"
