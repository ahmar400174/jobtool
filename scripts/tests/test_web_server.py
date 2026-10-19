"""Offline checks for web validation, isolated results, and Excel downloads."""
import sys
import time
from pathlib import Path
from unittest.mock import patch

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import web_server as server
from job_finder.aggregator import SourceCounts
from job_finder.models import JobListing


@pytest.fixture
def client():
    server.app.config["TESTING"] = True
    server.linkedin_results.clear()
    with server._lock:
        server._searches.clear()
        server._active = False
    yield server.app.test_client()


def payload(**changes):
    result = {"role": "Software Engineer", "country": "Worldwide", "start_date": "2026-10-01",
              "end_date": "2026-10-16", "sources": ["remotive"]}
    result.update(changes)
    return result


def test_health_and_web_page(client):
    assert client.get("/api/health").json["status"] == "ok"
    assert b"Job Finder" in client.get("/").data
    assert len(client.get("/api/config").json["sources"]) == 13


@pytest.mark.parametrize("changes", [
    {"role": ""}, {"country": "Invalid country"}, {"sources": []}, {"sources": ["unknown"]},
    {"sources": [{}]}, {"start_date": "bad"}, {"start_date": "2027-01-01"},
    {"end_date": "2028-01-01"}, {"remote_only": "false"},
])
def test_invalid_search_does_not_fetch(client, changes):
    with patch.object(server, "fetch_all") as fetch:
        assert client.post("/api/search", json=payload(**changes)).status_code == 400
    fetch.assert_not_called()


def test_server_does_not_accept_output_paths_or_unbounded_page_counts():
    settings = server.search_settings(payload(output={"directory": "/somewhere"}, max_pages=10000))
    assert settings["sources"]["linkedin"]["max_pages"] <= 3
    assert settings["sources"]["remotive"]["enabled"]
    assert not settings["sources"]["linkedin"]["enabled"]
    assert settings["output"]["directory"] != "/somewhere"


def test_cors_accepts_firebase_and_rejects_other_sites(client):
    accepted = client.options("/api/search", headers={"Origin": "https://jobs-tool.web.app"})
    assert accepted.status_code == 204
    assert accepted.headers["Access-Control-Allow-Origin"] == "https://jobs-tool.web.app"
    denied = client.post("/api/search", json=payload(), headers={"Origin": "https://example.org"})
    assert denied.status_code == 403
    assert "Access-Control-Allow-Origin" not in denied.headers


def test_running_search_returns_busy_without_creating_thread(client):
    server._active = True
    assert client.post("/api/search", json=payload()).status_code == 429


def test_async_search_can_be_polled_and_downloaded(client):
    job = JobListing("Software Engineer", "Acme", "Worldwide", "Remote", "2026-10-12", "https://example.org/job", "Remotive")
    counts = SourceCounts()
    counts.statuses = {"Remotive": "partial"}
    counts.warnings = {"Remotive": ["Recent feed only"]}

    def fetch(settings, on_progress):
        on_progress({"source": "Remotive", "status": "partial", "count": 1})
        return [job], counts

    with patch.object(server, "fetch_all", side_effect=fetch):
        response = client.post("/api/search", json=payload())
        assert response.status_code == 202
        search_id = response.json["id"]
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            data = client.get(f"/api/search/{search_id}").json
            if data["status"] != "running":
                break
            time.sleep(0.02)
    assert data["status"] == "complete"
    assert data["jobs"][0]["company"] == "Acme"
    assert data["progress"][0]["source"] == "Remotive"
    assert data["warnings"]["Remotive"] == ["Recent feed only"]
    assert "export" not in data and "created" not in data
    download = client.get(f"/api/search/{search_id}/excel")
    assert download.status_code == 200
    assert download.data.startswith(b"PK")
    assert not server._active
    assert client.get("/api/search/not-a-search").status_code == 404


def test_failed_search_releases_server(client):
    server._active = True
    server._searches["example"] = {"status": "running", "created": time.monotonic(), "progress": []}
    with patch.object(server, "fetch_all", side_effect=RuntimeError("internal failure")):
        server.run_search("example", server.search_settings(payload()))
    result = client.get("/api/search/example").json
    assert result["status"] == "failed"
    assert "internal failure" not in result["error"]
    assert not server._active


def test_expired_results_and_downloads_are_removed(client):
    server._searches["expired"] = {"status": "complete", "created": time.monotonic() - 3601, "export": b"file"}
    assert client.get("/api/search/expired").status_code == 404
    assert client.get("/api/search/expired/excel").status_code == 404


def test_stateless_source_returns_filtered_results_and_bounded_settings(client):
    counts = SourceCounts()
    counts.statuses = {"Remotive": "partial"}
    counts["Remotive"] = 5
    job = JobListing("Software Engineer", "Acme", "Worldwide", "Remote", "2026-10-12", "https://example.org/job", "Remotive")
    with patch.object(server, "fetch_all", return_value=([job], counts)) as fetch:
        response = client.post("/api/source", json=payload())
    assert response.status_code == 200
    assert response.json["jobs"][0]["company"] == "Acme"
    settings = fetch.call_args.args[0]
    assert settings["sources"]["linkedin"]["max_pages"] == 1
    assert settings["sources"]["hackernews"]["max_comments"] == 30
    assert not server._searches
    assert client.post("/api/source", json=payload(sources=["remotive", "linkedin"])).status_code == 400


def test_stateless_merge_and_export_keep_the_same_unique_jobs(client):
    import io
    from openpyxl import load_workbook
    first = JobListing("Software Engineer", "Acme", "Worldwide", "Remote", "2026-10-12", "https://example.org/remote", "Remotive")
    second = JobListing("Software Engineer", "Acme", "Worldwide", "Remote", "2026-10-12", "https://example.org/linkedin", "LinkedIn")
    rows = [first.to_row(), second.to_row()]
    merged = client.post("/api/merge", json={"jobs": rows})
    assert merged.status_code == 200
    assert merged.json["count"] == 1
    assert merged.json["jobs"][0]["source"] == "LinkedIn"
    exported = client.post("/api/export", json={"jobs": rows})
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.data), read_only=True)
    try:
        assert workbook["Jobs"].max_row == 2
        assert workbook["Jobs"]["F2"].value == "LinkedIn"
    finally:
        workbook.close()
    assert not server._searches


@pytest.mark.parametrize("rows", [[{}], [{"title": 123}], ["invalid"], [{"unexpected": "field"}]])
def test_stateless_export_rejects_invalid_job_records(client, rows):
    assert client.post("/api/merge", json={"jobs": rows}).status_code == 400
    assert client.post("/api/export", json={"jobs": rows}).status_code == 400


def test_serverless_mode_does_not_start_background_threads(client, monkeypatch):
    monkeypatch.setenv("VERCEL", "1")
    assert client.post("/api/search", json=payload()).status_code == 400
    assert not server._active


def test_linkedin_endpoint_reuses_successful_search(client):
    counts = SourceCounts()
    counts.statuses = {"LinkedIn": "ok"}
    counts["LinkedIn"] = 1
    job = JobListing("Software Engineer", "Acme", "Worldwide", "Remote", "2026-10-12", "https://www.linkedin.com/jobs/view/123456789", "LinkedIn")
    with patch.object(server, "fetch_all", return_value=([job], counts)) as fetch:
        first = client.post("/api/source", json=payload(sources=["linkedin"]))
        second = client.post("/api/source", json=payload(sources=["linkedin"]))
    assert first.status_code == second.status_code == 200
    assert not first.json["cache_info"]["cached"]
    assert second.json["cache_info"]["cached"]
    assert second.json["jobs"] == first.json["jobs"]
    fetch.assert_called_once()


def test_linkedin_endpoint_reports_in_flight_request(client):
    assert server.linkedin_results.request_lock.acquire(blocking=False)
    try:
        with patch.object(server, "fetch_all") as fetch:
            response = client.post("/api/source", json=payload(sources=["linkedin"]))
        assert response.status_code == 429
        assert response.headers["Retry-After"] == "10"
        assert response.json["retry_after_seconds"] == 10
        fetch.assert_not_called()
    finally:
        server.linkedin_results.request_lock.release()
